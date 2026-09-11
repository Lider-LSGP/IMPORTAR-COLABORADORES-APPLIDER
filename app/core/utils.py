"""
╔══════════════════════════════════════════════════════════════════════════╗
║  UTILS.PY — Funções utilitárias do Importador LiderLimp                  ║
╠══════════════════════════════════════════════════════════════════════════╣
║  ÍNDICE                                                                  ║
║  1. Normalização de textos                                               ║
║  2. Documentos / e-mail                                                  ║
║  3. Datas                                                                ║
║  4. Sexo / Supervisor / Ativo / Escolaridade                             ║
║  5. Escala de trabalho                                                   ║
║  6. Salário                                                              ║
║  7. Horário de trabalho                                                  ║
║  8. PCD (Pessoa com Deficiência)                                         ║
║  9. Grande Vitória - ES (VT / beneficioemconta)                          ║
║ 10. Conversão XLSX → CSV / DataFrame → XLSX                              ║
╚══════════════════════════════════════════════════════════════════════════╝
"""
from __future__ import annotations

import calendar
import io
import re
import unicodedata
from datetime import date, datetime

import pandas as pd
from openpyxl import load_workbook


# ══════════════════════════════════════════════════════════════════════════
# 1) NORMALIZAÇÃO DE TEXTOS
# ══════════════════════════════════════════════════════════════════════════

def normalizar_texto(valor) -> str:
    """
    Padroniza um texto para comparação segura:
      - maiúsculas
      - sem espaços nas pontas
      - SEM nenhum acento (Ã, Â, Á, À, É, Ê, Í, Ó, Õ, Ú, Ç ...)

    Isso garante que "BRITÂNIA" case com "BRITANIA",
    "VIGILÂNCIA" com "VIGILANCIA" e assim por diante.
    """
    if pd.isna(valor):
        return ""
    texto = str(valor).strip().upper()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return texto


# ══════════════════════════════════════════════════════════════════════════
# 2) DOCUMENTOS / E-MAIL
# ══════════════════════════════════════════════════════════════════════════

def limpar_documento(valor, tamanho: int | None = None) -> str:
    """Mantém zeros à esquerda e remove o `.0` inserido pelo Excel."""
    if pd.isna(valor) or str(valor).strip() == "" or str(valor).lower() == "nan":
        return ""
    s = str(valor).strip()
    if s.endswith(".0"):
        s = s[:-2]
    if tamanho:
        return s.zfill(tamanho)
    return s


def limpar_email(email) -> str:
    if pd.isna(email) or str(email).strip() == "":
        return ""
    email = str(email).strip().lower()
    email = "".join(c for c in unicodedata.normalize("NFKD", email) if not unicodedata.combining(c))
    return email


# ══════════════════════════════════════════════════════════════════════════
# 3) DATAS
# ══════════════════════════════════════════════════════════════════════════

def calcular_dias_entre(data_inicio, data_fim) -> str:
    """
    Diferença em DIAS entre duas datas (para os períodos de experiência).
    Retorna "" se alguma das datas for vazia/inválida.
    """
    try:
        # format="mixed" aceita tanto "01/09/2026" (BR) quanto "2026-09-01"
        # (ISO — formato que formatar_data devolve dentro do pipeline)
        d1 = pd.to_datetime(data_inicio, format="mixed", dayfirst=True, errors="coerce")
        d2 = pd.to_datetime(data_fim, format="mixed", dayfirst=True, errors="coerce")
        if pd.isna(d1) or pd.isna(d2):
            return ""
        return str((d2 - d1).days)
    except Exception:
        return ""


def formatar_data(valor) -> str | None:
    try:
        return pd.to_datetime(valor, dayfirst=True).strftime("%Y-%m-%d")
    except Exception:
        return None


def extrair_mes(valor):
    try:
        return pd.to_datetime(valor, dayfirst=True).strftime("%m")
    except Exception:
        return None


def extrair_ano(valor):
    try:
        return pd.to_datetime(valor, dayfirst=True).strftime("%Y")
    except Exception:
        return None


def extrair_mes_ano(valor):
    try:
        return pd.to_datetime(valor, dayfirst=True).strftime("%m/%Y")
    except Exception:
        return None


def calcular_idade(data_nascimento, referencia: datetime | None = None) -> float | None:
    """
    Idade exata (em anos, com fração) na data de referência (padrão: hoje).
    Retorna None se a data de nascimento for inválida/vazia.
    """
    try:
        nasc = pd.to_datetime(data_nascimento, dayfirst=True)
        ref = referencia or datetime.now()
        return (ref - nasc).days / 365.25
    except Exception:
        return None


def eh_menor_de_idade(data_nascimento, referencia: datetime | None = None) -> bool:
    """True se a pessoa tem MENOS de 18 anos na data de referência."""
    idade = calcular_idade(data_nascimento, referencia)
    return idade is not None and idade < 18


# ══════════════════════════════════════════════════════════════════════════
# 4) SEXO / SUPERVISOR / ATIVO / ESCOLARIDADE
# ══════════════════════════════════════════════════════════════════════════

def formatar_sexo(valor) -> str:
    v = str(valor).strip().upper()
    return "F" if v.startswith("F") else "M" if v.startswith("M") else ""


IDS_CARGOS_SUPERVISOR = [11, 12, 13, 14, 42]


def checar_supervisor(valor) -> str:
    try:
        if pd.isna(valor):
            return "Não"
        if int(float(valor)) in IDS_CARGOS_SUPERVISOR:
            return "Sim"
    except Exception:
        pass
    return "Não"


def checar_ativo(valor) -> str:
    v = str(valor).strip().lower()
    if "trabalhando" in v:
        return "Sim"
    if v in ["nan", "", "none", "nat", "01/01/1900", "00/00/0000"]:
        return "Sim"
    return "Não"


def escolher_escolaridade(row) -> str:
    sit = str(row.get("Situação", "")).strip().lower()
    grau = str(row.get("Grau instrução", "")).strip()
    if sit in ["", "nan", "none"] or "trabalhando" in sit:
        return grau
    return row.get("Situação", "")


# ══════════════════════════════════════════════════════════════════════════
# 5) ESCALA DE TRABALHO
# ══════════════════════════════════════════════════════════════════════════

def calcular_variacao_escala(data_adm_str) -> str:
    if pd.isna(data_adm_str) or str(data_adm_str).strip() == "":
        return "12x36"
    try:
        data_adm = pd.to_datetime(data_adm_str, dayfirst=True)
        hoje = datetime.now()
        paridade = "P" if data_adm.day % 2 == 0 else "I"
        ano, mes = data_adm.year, data_adm.month
        while (ano < hoje.year) or (ano == hoje.year and mes < hoje.month):
            ultimo_dia_mes = calendar.monthrange(ano, mes)[1]
            if ultimo_dia_mes in [31, 29]:
                paridade = "I" if paridade == "P" else "P"
            mes += 1
            if mes > 12:
                mes = 1
                ano += 1
        return f"12x36{paridade}"
    except Exception:
        return "12x36"


def identificar_tipo_escala(jornada, data_adm) -> str:
    if pd.isna(jornada):
        return calcular_variacao_escala(data_adm)
    j = str(jornada).upper().strip()
    if j in ["", "NAN", "NONE"]:
        return calcular_variacao_escala(data_adm)
    j_norm = re.sub(r"[^A-Z0-9]", "", j)
    if "12X36" in j_norm:
        return calcular_variacao_escala(data_adm)
    if "5X2" in j_norm or "5X02" in j_norm or "05X02" in j_norm:
        return "5x2"
    if "6X1" in j_norm or "6X01" in j_norm or "06X01" in j_norm:
        return "6x1"
    return ""


# ══════════════════════════════════════════════════════════════════════════
# 6) SALÁRIO
# ══════════════════════════════════════════════════════════════════════════

def formatar_salario(valor) -> str:
    if pd.isna(valor) or str(valor).strip() in ["", "nan", "NaN"]:
        return "0.00"
    v = str(valor).strip().replace(".", "").replace(",", "")
    if len(v) <= 4:
        # Provavelmente já está num formato simples como "1500"
        try:
            f = float(valor)
            return f"{f:.2f}"
        except Exception:
            return "0.00"
    v_cortado = v[:-4]
    parte_inteira = v_cortado[:-2] or "0"
    centavos = v_cortado[-2:].zfill(2)
    return f"{parte_inteira}.{centavos}"


# ══════════════════════════════════════════════════════════════════════════
# 7) HORÁRIO DE TRABALHO
# ══════════════════════════════════════════════════════════════════════════

def extrair_padrao_horario(texto):
    if pd.isna(texto):
        return None
    texto = str(texto).upper()
    escala = re.search(r"(\d{1,2})\s*X\s*(\d{1,2})", texto)
    if not escala:
        return None
    esc1 = escala.group(1).zfill(2)
    esc2 = escala.group(2).zfill(2)
    escala_formatada = f"{esc1}X{esc2}"
    texto_limpo = texto.replace("H", ":").replace("-", " ").replace("/", " ").replace("AS", " ")
    possiveis = re.findall(r"\d{1,2}:?\d{0,2}", texto_limpo)
    horarios_formatados = []
    for h in possiveis:
        if ":" in h:
            partes = h.split(":")
            if len(partes) == 2:
                hora = partes[0].zfill(2)
                minuto = partes[1].zfill(2) if partes[1] else "00"
                horarios_formatados.append(f"{hora}:{minuto}")
        else:
            if len(h) == 4:
                horarios_formatados.append(f"{h[:2]}:{h[2:]}")
            elif len(h) == 3:
                horarios_formatados.append(f"0{h[0]}:{h[1:]}")
            elif len(h) <= 2:
                horarios_formatados.append(f"{h.zfill(2)}:00")
    horarios_validos = []
    for h in horarios_formatados:
        try:
            hora, minuto = h.split(":")
            if int(hora) <= 23 and int(minuto) <= 59:
                horarios_validos.append(h)
        except Exception:
            continue
    if not horarios_validos:
        return None
    return escala_formatada + "_" + "_".join(horarios_validos)


def calcular_carga_horaria(horario_str) -> float:
    """
    Soma os intervalos de trabalho de um texto de jornada.
    Ex.: "07:30 11:00 12:00 17:30" → (11:00−07:30) + (17:30−12:00) = 9,0h

    CORREÇÃO v2: horários que viram o dia (ex.: 19:00 → 07:00) agora
    somam 24h em vez de dar valor negativo.
    """
    try:
        texto = str(horario_str)
        horarios = re.findall(r"\d{2}:\d{2}", texto)
        total_horas = 0.0
        for i in range(0, len(horarios), 2):
            if i + 1 < len(horarios):
                h1, m1 = map(int, horarios[i].split(":"))
                h2, m2 = map(int, horarios[i + 1].split(":"))
                entrada = h1 + m1 / 60
                saida = h2 + m2 / 60
                if saida <= entrada:      # vira o dia (turno noturno)
                    saida += 24
                total_horas += (saida - entrada)
        return total_horas
    except Exception:
        return 0.0


# ══════════════════════════════════════════════════════════════════════════
# 8) PCD (PESSOA COM DEFICIÊNCIA)
# ══════════════════════════════════════════════════════════════════════════

def definir_tipo_pcd(row) -> str:
    possui = str(row.get("Possui deficiência", "")).strip().upper()
    if possui != "SIM":
        return ""
    fisica = str(row.get("Deficiência física", "")).strip().upper() == "SIM"
    visual = str(row.get("Deficiência visual", "")).strip().upper() == "SIM"
    auditiva = str(row.get("Deficiência auditiva", "")).strip().upper() == "SIM"
    intelectual = str(row.get("Deficiência intelectual", "")).strip().upper() == "SIM"
    mental = str(row.get("Deficiência mental", "")).strip().upper() == "SIM"
    tipos = []
    if fisica: tipos.append("Física")
    if visual: tipos.append("Visual")
    if auditiva: tipos.append("Auditiva")
    if intelectual: tipos.append("Intelectual")
    if mental: tipos.append("Psicossocial ou por Saúde Mental")
    if len(tipos) > 1:
        return "Múltipla"
    if len(tipos) == 1:
        return tipos[0]
    return ""


# ══════════════════════════════════════════════════════════════════════════
# 9) GRANDE VITÓRIA - ES  (VT 58/59  e  coluna beneficioemconta)
# ══════════════════════════════════════════════════════════════════════════
# Área atendida pelo GVBUS (Transcol): Vitória, Vila Velha, Serra,
# Cariacica e Viana.  Validado com o setor em set/2026.
#
# ⚠️ A lista pode ser editada sem mexer no código, na aba LISTAS da
#    planilha do Google Sheets (lista CIDADES_GRANDE_VITORIA). Os valores
#    abaixo são o PADRÃO usado quando a planilha não está disponível.

CIDADES_GRANDE_VITORIA = [
    "VITORIA",
    "VILA VELHA",
    "SERRA",
    "CARIACICA",
    "VIANA",
]


def cidade_eh_grande_vitoria(cidade, cidades_gv: list | None = None) -> bool:
    """
    True se a cidade informada faz parte da Grande Vitória - ES.
    Compara o texto normalizado (sem acento) contra a lista oficial.
    Funciona tanto para "VILA VELHA" (Domínio) quanto para "Vila Velha".

    `cidades_gv`: lista customizada (aba LISTAS do Google Sheets);
    se omitida, usa o padrão CIDADES_GRANDE_VITORIA.
    """
    if pd.isna(cidade):
        return False
    lista = cidades_gv if cidades_gv else CIDADES_GRANDE_VITORIA
    return normalizar_texto(cidade) in lista


def definir_beneficio_em_conta(cidade, cidades_gv: list | None = None) -> str:
    """
    Regra da coluna `beneficioemconta` (novo layout de importação):
      - mora NA Grande Vitória  → "Não"  (recebe VT em cartão/passe)
      - mora FORA (interior/outros estados) → "Sim" (benefício em conta)
      - cidade vazia/desconhecida → "Sim" (mais seguro; a linha também é
        destacada em vermelho para revisão manual)
    """
    if pd.isna(cidade) or normalizar_texto(cidade) == "":
        return "Sim"
    return "Não" if cidade_eh_grande_vitoria(cidade, cidades_gv) else "Sim"


# ══════════════════════════════════════════════════════════════════════════
# 10) CONVERSÃO DE ARQUIVOS (XLSX → CSV  /  DataFrame → XLSX)
# ══════════════════════════════════════════════════════════════════════════

def _valor_exato_da_celula(cell) -> str:
    valor = cell.value
    if valor is None:
        return ""
    if isinstance(valor, str):
        return valor
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%Y-%m-%d")
    formato = (cell.number_format or "").split(";")[0].strip()
    if re.fullmatch(r"0+", formato):
        try:
            return f"{int(valor):0{len(formato)}d}"
        except Exception:
            return str(valor)
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


def xlsx_bytes_para_csv_bytes(xlsx_bytes: bytes) -> bytes:
    """Converte um .xlsx em bytes para um CSV em bytes preservando texto."""
    import csv
    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    ws = wb.active
    buf = io.StringIO()
    writer = csv.writer(buf)
    for row in ws.iter_rows():
        writer.writerow([_valor_exato_da_celula(cell) for cell in row])
    return buf.getvalue().encode("utf-8")


def dataframe_para_xlsx_bytes(df: pd.DataFrame, sheet_name: str = "Sheet1") -> bytes:
    """Salva DataFrame em xlsx em memória."""
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name=sheet_name)
    return out.getvalue()
