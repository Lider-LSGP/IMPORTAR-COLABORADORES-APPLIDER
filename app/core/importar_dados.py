"""
╔══════════════════════════════════════════════════════════════════════════╗
║  IMPORTAR_DADOS.PY — Geração da planilha "Importar Layout.xlsx"          ║
║  (importação de colaboradores Domínio → AppLider/EasyApp)                ║
╠══════════════════════════════════════════════════════════════════════════╣
║  ÍNDICE                                                                  ║
║  1. LAYOUT_COLUNAS — ordem OFICIAL das colunas (fixa no código)          ║
║  2. Regra de POSTO ADM (sede) por empresa                                ║
║  3. Regra de MENOR APRENDIZ (< 18 anos)                                  ║
║  4. processar_layout — montagem campo a campo                            ║
║  5. Diagnóstico de campos críticos (destaque vermelho)                   ║
║  6. aplicar_destaque_vermelho                                            ║
╠══════════════════════════════════════════════════════════════════════════╣
║  NOVIDADES v2 (set/2026):                                                ║
║  • Colunas não dependem mais da planilha "Colunas Originais.xlsx":       ║
║    a ordem está FIXA no código (constante LAYOUT_COLUNAS).               ║
║  • Nova coluna `beneficioemconta` (substitui `nivelcoberturaposto`).     ║
║  • `cpf` renomeada para `cnpj_cpf` (mesma informação de CPF).            ║
║  • `nomeescala` e `escalatrabalho_id` trocaram de posição.               ║
║  • Postos "ADM - ..." / "BASE - ..." são vinculados ao posto ADM         ║
║    correto de acordo com a EMPRESA do colaborador.                       ║
║  • Menores de 18 anos recebem a função 76 — MENOR APRENDIZ.              ║
╚══════════════════════════════════════════════════════════════════════════╝
"""
from __future__ import annotations

import io
import re
from datetime import datetime
from typing import Tuple, List, Dict

import pandas as pd


def _ler_excel_auto(dados, **kwargs):
    """Lê bytes de Excel detectando o formato pelo conteúdo
    (xlsx = assinatura PK/zip; caso contrário assume .xls)."""
    buf = io.BytesIO(dados)
    head = buf.read(4)
    buf.seek(0)
    if head[:2] == b"PK":
        return pd.read_excel(buf, engine="openpyxl", **kwargs)
    return pd.read_excel(buf, engine="xlrd", **kwargs)

from .utils import (
    formatar_sexo, limpar_documento, limpar_email, normalizar_texto,
    formatar_data, formatar_salario, extrair_mes, extrair_ano,
    extrair_mes_ano, definir_tipo_pcd, identificar_tipo_escala,
    extrair_padrao_horario, escolher_escolaridade, checar_supervisor,
    checar_ativo, eh_menor_de_idade, definir_beneficio_em_conta,
    calcular_dias_entre,
)
from .mapas import carregar_todos_os_mapas
from .config import RegrasConfig, PADRAO_EMPRESA_PARA_POSTO_ADM


# ══════════════════════════════════════════════════════════════════════════
# 1) LAYOUT OFICIAL DE IMPORTAÇÃO (ordem das colunas — NÃO reordenar!)
#    Espelha o "Colunas Originais (ATUALIZADO).xlsx" de set/2026.
#    Mudanças em relação ao layout antigo:
#      - posição 7 : "cpf"  →  "cnpj_cpf"
#      - final     : "escalatrabalho_id" ↔ "nomeescala" trocaram de lugar
#      - última    : "nivelcoberturaposto"  →  "beneficioemconta"
# ══════════════════════════════════════════════════════════════════════════
LAYOUT_COLUNAS = [
    "id", "ativo", "foto", "tipocontrolepessoa", "tipojuridico",
    "sexo", "matricula", "cnpj_cpf", "cnpj", "nome",
    "nomefantasia", "tipofuncionario", "empresa_id", "postotrabalho_id", "cep",
    "codigoibgemunicipio", "codigoibgeestado", "cidade_id", "tipoendereco", "rua",
    "cidadeparceiro", "numero", "bairro", "uf", "complemento",
    "telefone", "celular", "email", "ci", "cidataemissao",
    "ciemissor", "pis", "pisdata", "nomeconjuge", "cpfconjuge",
    "ciconjunge", "tituloeleitor", "tituloeleitorsecao", "tituloeleitorzona", "certificadomilitar",
    "cnh", "cnhcategoria", "cnhdataemissao", "cnhdatavalidade", "cnhuf",
    "ctpsnumero", "ctpsserie", "ctpsuf", "ctpsdigital", "ctpsdataemissao",
    "estadocivil_id", "municipios_id", "nacionalidade_id", "escolaridade_id", "datanascimento",
    "racaetinia_id", "paises_id", "tipopcd", "nomepai", "nomemae",
    "dataadmissao", "tipoadmissao", "regimevinculo", "cidadetrabalho", "planosaude",
    "planoodonto", "possuisegurovida", "salariobase", "tiposalario", "tipirecebimento",
    "periodo1diasexp", "datavenci1diasexp", "periodo2diasexp", "datavenci2diasexp", "website",
    "codigoentidade", "nomepresidente", "nomesecretario", "databasesindicato", "obs",
    "mes", "ano", "mes_ano", "nomeempresa", "nomepostotrabalho",
    "funcao_id", "sindicato_id", "created_at", "updated_at", "delete_at",
    "horariotrabalho_id", "nomehorariotrabalho", "supervisor", "usoregcreated_at", "usoregupdated_at",
    "usoregdelete_at", "datademissao", "motivodesligamento", "situacaofuncionario", "datacadastro",
    "nomeescala", "tipoescala", "escalatrabalho_id", "nomefuncao", "nomesindicato",
    "beneficioemconta",
]


# Campos críticos que, se vazios/None, ficam em VERMELHO na planilha.
CAMPOS_CRITICOS_VERMELHO = [
    "postotrabalho_id",   # posto de serviço
    "nomepostotrabalho",
    "funcao_id",          # função
    "nomefuncao",
    "escalatrabalho_id",  # escala de trabalho
    "tipoescala",
    "nomeescala",
    "horariotrabalho_id",
    "nomehorariotrabalho",
]


# ══════════════════════════════════════════════════════════════════════════
# 2) REGRA DE POSTO "ADM" (SEDE) — vinculado à EMPRESA do colaborador
# ══════════════════════════════════════════════════════════════════════════
# Quando o posto vindo do Domínio é um posto administrativo genérico
# ("ADM - SEDE", "BASE - ADMINISTRATIVO", "VSP BASE", ...), o posto correto
# no EasyApp depende da EMPRESA em que a pessoa está sendo cadastrada:
#
#   Cód Emp Domínio  Empresa                                          Posto EasyApp
#   1                VSP VIGILANCIA E SEGURANCA PATRIMONIAL LTDA   →  ADM - VSP
#   2                ATIVA TERCEIRIZACAO DE MAO DE OBRA LTDA       →  ADM - ATIVA
#   3                LIDER MULTISSERVICOS LTDA                     →  ADM - LIDER MULTISSERVICOS
#   4                LIDER LIMPE LIMPEZA COMERCIAL LTDA            →  ADM - LIDER LIMPE
#   5                B2WE BUSINESS TO WE ASSESSORIA EMPRESARIAL    →  ADM - B2WE
#
# (validado com o setor em set/2026)
#
# ⚠️ Os valores abaixo são os nomes no DOMÍNIO (chave de busca na aba POSTOS
#    do Mapeamento Sistema). No EasyApp eles viram: ADM - VSP, ADM - ATIVA,
#    ADM - LIDER MULTISSERVICOS, ADM - LIDER LIMPE e ADM - B2WE.

# ⚠️ Este mapa agora pode ser EDITADO sem mexer no código, na aba LISTAS
#    da planilha do Google Sheets (lista EMPRESA_PARA_POSTO_ADM).
#    Os valores abaixo são o PADRÃO (espelho da aba LISTAS em set/2026).
EMPRESA_PARA_POSTO_ADM = PADRAO_EMPRESA_PARA_POSTO_ADM

# Postos que disparam a regra: começam com "ADM" ou "BASE"
# (ex.: "ADM - SEDE", "ADM - ATIVA", "BASE - ADMINISTRATIVO")
# + alguns nomes conhecidos que não começam assim:
POSTOS_ADM_EXTRAS = {"VSP BASE", "VSP - BASE", "BASE - VSP"}

# ⚠️ ATENÇÃO: postos operacionais com a palavra "BASE" no MEIO do nome
# (ex.: "POLICIA CIVIL - BASE AMARELA", "PMV SEMSU - ROMU/BASE",
# "SEACREST PETROLEO - BASE") NÃO entram nesta regra — por isso o teste
# é só no INÍCIO do texto + lista de exceções acima.

_REGEX_POSTO_ADM = re.compile(r"^(ADM|BASE)[\s\-/]")


def posto_eh_administrativo(posto_norm: str) -> bool:
    """True se o posto (já normalizado) é um posto administrativo/sede."""
    if not posto_norm:
        return False
    if _REGEX_POSTO_ADM.match(posto_norm):
        return True
    return posto_norm in POSTOS_ADM_EXTRAS


def resolver_posto_adm(cod_emp_dominio: str, mapa_empresa: dict | None = None) -> str | None:
    """Devolve o nome do posto ADM correto para a empresa (ou None).

    `mapa_empresa`: mapa customizado (aba LISTAS do Google Sheets);
    se omitido, usa o padrão EMPRESA_PARA_POSTO_ADM.
    """
    chave = str(cod_emp_dominio).strip()
    if chave.endswith(".0"):
        chave = chave[:-2]
    return (mapa_empresa or EMPRESA_PARA_POSTO_ADM).get(chave)


# ══════════════════════════════════════════════════════════════════════════
# 3) REGRA DE MENOR APRENDIZ
# ══════════════════════════════════════════════════════════════════════════
# Se a pessoa tem MENOS de 18 anos (pela data de nascimento), a função é
# substituída por "MENOR APRENDIZ" (ID 76 no EasyApp), independentemente
# do cargo vindo do Domínio. Os benefícios tratam essa regra no módulo
# de benefícios (menor aprendiz NÃO recebe VA, mas recebe extras e VT).

FUNCAO_MENOR_APRENDIZ_ID = 76
FUNCAO_MENOR_APRENDIZ_NOME = "MENOR APRENDIZ"


# ══════════════════════════════════════════════════════════════════════════
# 4) PROCESSAMENTO PRINCIPAL
# ══════════════════════════════════════════════════════════════════════════
def processar_layout(
    dominio_bytes: bytes,
    mapeamento_bytes: bytes,
    modelo_bytes: bytes | None = None,   # legado — não é mais necessário
    regras: RegrasConfig | None = None,  # listas editáveis (Google Sheets)
) -> Tuple[pd.DataFrame, List[Dict]]:
    """
    Lê Arquivo Domínio + Mapeamento e devolve:
      - DataFrame final no layout oficial do AppLider (LAYOUT_COLUNAS);
      - Lista de problemas (1 dict por linha) com campos críticos vazios.

    O parâmetro `modelo_bytes` (antiga planilha "Colunas Originais") é
    aceito por compatibilidade, mas IGNORADO: a ordem das colunas agora
    é fixa no código (constante LAYOUT_COLUNAS).

    `regras`: listas configuráveis (posto ADM por empresa, cidades da
    Grande Vitória) vindas da aba LISTAS do Google Sheets; se omitidas,
    valem os padrões do código.
    """
    regras = regras or RegrasConfig()
    # ─── 4.1 Entradas ──────────────────────────────────────────────────
    arq_dom = _ler_excel_auto(dominio_bytes, dtype=str)
    mapa = carregar_todos_os_mapas(mapeamento_bytes)

    arq_lider = pd.DataFrame(columns=LAYOUT_COLUNAS)
    agora = datetime.now()

    def clean(col):
        """Coluna do Domínio em MAIÚSCULAS e sem espaços nas pontas."""
        if col in arq_dom.columns:
            return arq_dom[col].astype(str).str.upper().str.strip()
        return pd.Series([""] * len(arq_dom))

    # ─── 4.2 Identificação / Empresa ───────────────────────────────────
    arq_lider["empresa_id"]  = clean("Cód Emp").map(mapa.get("MAPA_EMPRESAS_ID"))
    arq_lider["nomeempresa"] = clean("Cód Emp").map(mapa.get("MAPA_EMPRESAS_NOME"))
    arq_lider["sexo"]        = arq_dom.get("Sexo", pd.Series([""] * len(arq_dom))).apply(formatar_sexo)
    arq_lider["matricula"]   = arq_dom.get("Cód eSocial")
    arq_lider["nome"]        = arq_dom.get("Nome")
    arq_lider["cnpj_cpf"]    = arq_dom.get("CPF", pd.Series([""] * len(arq_dom))).apply(lambda x: limpar_documento(x, 11))

    # ─── 4.3 Dados pessoais ────────────────────────────────────────────
    arq_lider["racaetinia_id"] = arq_dom.get("Raça/Cor", pd.Series([""] * len(arq_dom))).apply(normalizar_texto).map(mapa.get("MAPA_RACA_ID"))

    escolaridade_escolhida = arq_dom.apply(escolher_escolaridade, axis=1)
    escolaridade_norm = escolaridade_escolhida.apply(normalizar_texto)
    arq_lider["escolaridade_id"] = escolaridade_norm.map(mapa.get("MAPA_ESCOLARIDADE_ID"))

    ec_norm = arq_dom.get("Estado Civil", pd.Series([""] * len(arq_dom))).apply(normalizar_texto)
    arq_lider["estadocivil_id"] = ec_norm.map(mapa.get("MAPA_ESTADOCIVIL_ID"))

    arq_lider["telefone"] = arq_dom.get("Telefone", pd.Series([""] * len(arq_dom))).apply(limpar_documento)
    arq_lider["celular"]  = arq_dom.get("Celular",  pd.Series([""] * len(arq_dom))).apply(limpar_documento)
    arq_lider["email"]    = arq_dom.get("Email",    pd.Series([""] * len(arq_dom))).apply(limpar_email)
    arq_lider["nomepai"]  = arq_dom.get("Nome Pai")
    arq_lider["nomemae"]  = arq_dom.get("Nome Mãe")

    # ─── 4.4 Documentos ────────────────────────────────────────────────
    arq_lider["ci"]                 = arq_dom.get("RG")
    arq_lider["ciemissor"]          = arq_dom.get("Orgão RG")
    arq_lider["pis"]                = arq_dom.get("PIS", pd.Series([""] * len(arq_dom))).apply(limpar_documento)
    arq_lider["tituloeleitor"]      = arq_dom.get("Titulo", pd.Series([""] * len(arq_dom))).apply(limpar_documento)
    arq_lider["tituloeleitorsecao"] = arq_dom.get("Seção", pd.Series([""] * len(arq_dom))).apply(limpar_documento)
    arq_lider["tituloeleitorzona"]  = arq_dom.get("Zona",  pd.Series([""] * len(arq_dom))).apply(limpar_documento)
    arq_lider["certificadomilitar"] = arq_dom.get("Reservista", pd.Series([""] * len(arq_dom))).apply(limpar_documento)
    arq_lider["cnh"]                = arq_dom.get("CNH", pd.Series([""] * len(arq_dom))).apply(limpar_documento)
    arq_lider["cnhcategoria"]       = arq_dom.get("Categoria CNH", pd.Series([""] * len(arq_dom))).astype(str).str.strip().str.upper()
    arq_lider["ctpsnumero"]         = arq_dom.get("CTPS", pd.Series([""] * len(arq_dom))).apply(limpar_documento)
    arq_lider["ctpsserie"]          = arq_dom.get("Serie CTPS", pd.Series([""] * len(arq_dom))).apply(limpar_documento)
    arq_lider["ctpsuf"]             = arq_dom.get("UF CTPS")
    arq_lider["tipopcd"]            = arq_dom.apply(definir_tipo_pcd, axis=1)

    # ─── 4.5 Endereço + beneficioemconta ───────────────────────────────
    #   beneficioemconta (NOVO): definido pela CIDADE DE RESIDÊNCIA
    #     • Grande Vitória-ES (Vitória, Vila Velha, Serra, Cariacica,
    #       Viana)           → "Não"
    #     • Interior / fora do ES / cidade não informada → "Sim"
    #   (validado com o setor: decidir sempre pela cidade, não pelo CEP)
    cidade_bruta = arq_dom.get("Cidade", pd.Series([""] * len(arq_dom)))
    cid_normalizada = cidade_bruta.apply(normalizar_texto)
    arq_lider["cidade_id"]      = cid_normalizada.map(mapa.get("MAPA_CIDADES_ID"))
    arq_lider["municipios_id"]  = arq_lider["cidade_id"].map(mapa.get("MAPA_CIDADE_PARA_MUNICIPIO"))
    arq_lider["cidadeparceiro"] = cid_normalizada.map(mapa.get("MAPA_CIDADES_NOME"))
    arq_lider["cep"]            = arq_dom.get("Cep", pd.Series([""] * len(arq_dom))).apply(limpar_documento)
    arq_lider["rua"]            = arq_dom.get("Endereço")
    arq_lider["numero"]         = arq_dom.get("Numero", pd.Series([""] * len(arq_dom))).apply(limpar_documento)
    arq_lider["bairro"]         = arq_dom.get("Bairro")
    arq_lider["uf"]             = arq_dom.get("UF End")
    arq_lider["complemento"]    = arq_dom.get("Complemento")
    arq_lider["cidadetrabalho"] = cid_normalizada.map(mapa.get("MAPA_CIDADES_NOME"))
    arq_lider["beneficioemconta"] = cidade_bruta.apply(
        lambda c: definir_beneficio_em_conta(c, regras.cidades_grande_vitoria)
    )

    # ─── 4.6 Cargo, Posto, Sindicato ───────────────────────────────────
    posto_bruto   = arq_dom.get("Descrição Ccusto", pd.Series([""] * len(arq_dom)))
    posto_norm    = posto_bruto.apply(normalizar_texto)
    funcao_norm   = arq_dom.get("Descrição cargo",  pd.Series([""] * len(arq_dom))).apply(normalizar_texto)
    sind_norm     = arq_dom.get("Sindicato",        pd.Series([""] * len(arq_dom))).apply(normalizar_texto)
    cod_emp_clean = clean("Cód Emp")

    # 4.6a — Regra do POSTO ADM: se o posto é administrativo/sede,
    #        troca pelo posto ADM da empresa ANTES de consultar o mapa.
    posto_resolvido = posto_norm.copy()
    for i in range(len(posto_resolvido)):
        if posto_eh_administrativo(posto_resolvido.iloc[i]):
            posto_adm = resolver_posto_adm(cod_emp_clean.iloc[i], regras.empresa_para_posto_adm)
            if posto_adm:
                posto_resolvido.iloc[i] = normalizar_texto(posto_adm)

    arq_lider["postotrabalho_id"]  = posto_resolvido.map(mapa.get("MAPA_POSTOS_ID"))
    arq_lider["nomepostotrabalho"] = posto_resolvido.map(mapa.get("MAPA_POSTOS_NOME"))

    # 4.6b — Regra do MENOR APRENDIZ: < 18 anos → função 76.
    nascimento = arq_dom.get("Data nascimento", pd.Series([""] * len(arq_dom)))
    menores = nascimento.apply(eh_menor_de_idade)

    arq_lider["funcao_id"]  = funcao_norm.map(mapa.get("MAPA_FUNCOES_ID"))
    arq_lider["nomefuncao"] = funcao_norm.map(mapa.get("MAPA_FUNCOES_NOME"))
    arq_lider.loc[menores, "funcao_id"]  = FUNCAO_MENOR_APRENDIZ_ID
    arq_lider.loc[menores, "nomefuncao"] = FUNCAO_MENOR_APRENDIZ_NOME

    arq_lider["sindicato_id"]  = sind_norm.map(mapa.get("MAPA_SINDICATOS_ID"))
    arq_lider["nomesindicato"] = sind_norm.map(mapa.get("MAPA_SINDICATOS_NOME"))
    arq_lider["supervisor"]    = arq_lider["funcao_id"].apply(checar_supervisor)

    # ─── 4.7 Escala e horário ──────────────────────────────────────────
    arq_lider["tipoescala"] = arq_dom.apply(
        lambda r: identificar_tipo_escala(r.get("Jornada"), r.get("Admissão")), axis=1
    )
    escala_norm = arq_lider["tipoescala"].apply(normalizar_texto)
    arq_lider["escalatrabalho_id"] = escala_norm.map(mapa.get("MAPA_ESCALAS_ID"))
    arq_lider["nomeescala"] = "GRADE DA ESCALA " + arq_lider["tipoescala"].fillna("").astype(str)

    jornada_key = arq_dom.get("Jornada", pd.Series([""] * len(arq_dom))).apply(extrair_padrao_horario)
    arq_lider["horariotrabalho_id"]  = jornada_key.map(mapa.get("MAPA_HORARIOS_ID"))
    arq_lider["nomehorariotrabalho"] = jornada_key.map(mapa.get("MAPA_HORARIOS_NOME"))

    arq_lider["mes"]     = arq_dom.get("Admissão", pd.Series([""] * len(arq_dom))).apply(extrair_mes)
    arq_lider["ano"]     = arq_dom.get("Admissão", pd.Series([""] * len(arq_dom))).apply(extrair_ano)
    arq_lider["mes_ano"] = arq_dom.get("Admissão", pd.Series([""] * len(arq_dom))).apply(extrair_mes_ano)

    # ─── 4.8 Datas e salário ───────────────────────────────────────────
    arq_lider["dataadmissao"]      = arq_dom.get("Admissão", pd.Series([""] * len(arq_dom))).apply(formatar_data)
    arq_lider["cidataemissao"]     = arq_dom.get("Data EX", pd.Series([""] * len(arq_dom))).apply(formatar_data)
    arq_lider["cnhdataemissao"]    = arq_dom.get("Expedição CNH", pd.Series([""] * len(arq_dom))).apply(formatar_data)
    arq_lider["cnhdatavalidade"]   = arq_dom.get("Vencimento", pd.Series([""] * len(arq_dom))).apply(formatar_data)
    arq_lider["datanascimento"]    = nascimento.apply(formatar_data)
    arq_lider["datavenci1diasexp"] = arq_dom.get("Fim Determinado", pd.Series([""] * len(arq_dom))).apply(formatar_data)
    arq_lider["datavenci2diasexp"] = arq_dom.get("Fim Prorrogação", pd.Series([""] * len(arq_dom))).apply(formatar_data)
    arq_lider["salariobase"]       = arq_dom.get("Salário", pd.Series([""] * len(arq_dom))).apply(formatar_salario)

    # ─── Períodos de experiência (calculados por diferença de dias) ──────
    #   periodo1diasexp = datavenci1diasexp − dataadmissao
    #   periodo2diasexp = datavenci2diasexp − datavenci1diasexp
    #   (antes eram fixos em "45"; se faltar alguma data, fica vazio)
    arq_lider["periodo1diasexp"] = [
        calcular_dias_entre(a, v1)
        for a, v1 in zip(arq_lider["dataadmissao"], arq_lider["datavenci1diasexp"])
    ]
    arq_lider["periodo2diasexp"] = [
        calcular_dias_entre(v1, v2)
        for v1, v2 in zip(arq_lider["datavenci1diasexp"], arq_lider["datavenci2diasexp"])
    ]

    categoria_norm = arq_dom.get("Categoria", pd.Series([""] * len(arq_dom))).apply(normalizar_texto)
    arq_lider["tiposalario"] = categoria_norm.map(mapa.get("MAPA_TIPOSALARIO_NOME"))

    # ─── 4.9 Desligamento ──────────────────────────────────────────────
    # ─── Demissão: status / data / motivo ─────────────────────────────
    # O Domínio pode exportar as colunas "Data Demissão" e
    # "Motivo Demissão" TROCADAS (numa vem o status "Trabalhando"/
    # "Demitido" e na outra vem a DATA real da demissão). Aqui
    # detectamos automaticamente em qual coluna está a data e
    # preenchemos datademissao sempre que a pessoa estiver demitida.
    col_status_d = arq_dom.get("Data Demissão",   pd.Series([""] * len(arq_dom)))
    col_motivo_d = arq_dom.get("Motivo Demissão", pd.Series([""] * len(arq_dom)))
    mapa_desl_nome = mapa.get("MAPA_DESLIGAMENTO_NOME", {})

    _ativos, _datas_dem, _motivos_dem = [], [], []
    for _sv, _mv in zip(col_status_d, col_motivo_d):
        _d_status = formatar_data(_sv)   # data estava na coluna "Data Demissão"?
        _d_motivo = formatar_data(_mv)   # ou veio trocada p/ "Motivo Demissão"?
        _data_dem = _d_status or _d_motivo
        _status_norm = normalizar_texto(_sv)
        # Demitido se: status contém "DEMIT" OU existe data de demissão
        _demitido = ("DEMIT" in _status_norm) or (_data_dem is not None)
        _ativos.append("Não" if _demitido else "Sim")
        _datas_dem.append(_data_dem if _demitido else None)
        # Motivo: só usa o texto se ele NÃO for uma data (coluna trocada)
        _motivo_txt = None if _d_motivo else _mv
        _motivos_dem.append(
            mapa_desl_nome.get(normalizar_texto(_motivo_txt))
            if (_demitido and _motivo_txt is not None and str(_motivo_txt).strip() != "")
            else None
        )

    arq_lider["ativo"]              = _ativos
    arq_lider["datademissao"]       = _datas_dem
    arq_lider["motivodesligamento"] = _motivos_dem

    # ─── 4.10 Campos fixos ─────────────────────────────────────────────
    arq_lider["tipoadmissao"]         = "CLT"
    arq_lider["tipocontrolepessoa"]   = "Funcionário"
    arq_lider["tipofuncionario"]      = "Funcionário"
    arq_lider["ctpsdigital"]          = "Não"
    arq_lider["nacionalidade_id"]     = "1"
    arq_lider["paises_id"]            = "1"
    arq_lider["planosaude"]           = "Sim"
    arq_lider["planoodonto"]          = "Sim"
    arq_lider["possuisegurovida"]     = "Sim"
    arq_lider["tipirecebimento"]      = "Banco"
    arq_lider["tipoendereco"]         = "Rua"
    arq_lider["databasesindicato"]    = f"{agora.year}-01-01"
    arq_lider["datacadastro"]         = agora.strftime("%Y-%m-%d")

    # ══════════════════════════════════════════════════════════════════
    # 5) DIAGNÓSTICO DE CAMPOS CRÍTICOS VAZIOS (destaque vermelho)
    # ══════════════════════════════════════════════════════════════════
    problemas: List[Dict] = []
    for i in range(len(arq_lider)):
        vazios = []
        for campo in CAMPOS_CRITICOS_VERMELHO:
            val = arq_lider.iloc[i].get(campo) if campo in arq_lider.columns else None
            if pd.isna(val) or (isinstance(val, str) and val.strip() == "") or val is None:
                vazios.append(campo)
        if vazios:
            problemas.append({
                "linha_excel": i + 2,  # +1 cabeçalho, +1 base 1
                "matricula": arq_lider.iloc[i].get("matricula"),
                "nome": arq_lider.iloc[i].get("nome"),
                "campos_faltantes": vazios,
            })

    return arq_lider, problemas


# ══════════════════════════════════════════════════════════════════════════
# 6) DESTAQUE VERMELHO NA PLANILHA GERADA
# ══════════════════════════════════════════════════════════════════════════
def aplicar_destaque_vermelho(xlsx_bytes: bytes, problemas: List[Dict]) -> bytes:
    """
    Recebe o xlsx em bytes e marca as células dos campos faltantes em vermelho.
    Também escreve "FALTANDO" nas células vazias.
    """
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill, Font

    if not problemas:
        return xlsx_bytes

    wb = load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active

    # cabeçalho → posição da coluna
    header = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}

    fill_red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    font_red = Font(color="9C0006", bold=True)

    for p in problemas:
        linha = p["linha_excel"]
        for campo in p["campos_faltantes"]:
            col = header.get(campo)
            if col:
                c = ws.cell(row=linha, column=col)
                c.fill = fill_red
                c.font = font_red
                if c.value in (None, "", "nan"):
                    c.value = "FALTANDO"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
