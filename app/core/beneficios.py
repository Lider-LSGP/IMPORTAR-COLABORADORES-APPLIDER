"""
╔══════════════════════════════════════════════════════════════════════════╗
║  BENEFICIOS.PY — Motor de regras de benefícios (v4 — set/2026)           ║
║  Geração da planilha "Importar Beneficios.xlsx"                          ║
╠══════════════════════════════════════════════════════════════════════════╣
║  FONTE DAS REGRAS: aba CONEXAO da planilha de benefícios.                ║
║  As regras abaixo estão FIXAS no código (a planilha é a documentação).   ║
║                                                                          ║
║  REGRA DE OURO (validada com o setor):                                   ║
║    • Sigla VA  → benefício PRINCIPAL (só UM por colaborador)             ║
║    • Sigla CB / CF / VT → EXTRAS (acumulam, pode ter vários)             ║
║    • O match do posto/função segue EXATAMENTE o texto da CONEXAO.        ║
║                                                                          ║
║  ORDEM DE PRIORIDADE DO VA PRINCIPAL (a 1ª que bater vence):             ║
║    P0  Menor aprendiz (<18) → SEM VA                                     ║
║    P1  SUPERVISOR ............. 63  VA ADM 700        [função]           ║
║    P2  MOTOBOY ................ 57  VA MOTOBOY        [função]           ║
║    P3  MOTORISTA .............. 40  VA MOTORISTA      [função]           ║
║    P4  VIGILANTE .............. 56  VA VIGILANTE VSP  [função]           ║
║    P5  COZINHA do AMBIENTAL VALE  50  VA COZINHA VALE [função+posto]     ║
║    P6  VALE SA ................ 136 VA VALE           [posto]            ║
║    P7  VPORTS ................. 54  VA VPORTS         [posto]            ║
║    P8  MULTILIFT .............. 54  VA VPORTS         [posto]            ║
║    P9  IFES (começa com) ...... 52  VA IFE            [posto]            ║
║    P10 BRASILIA (contém) ...... 49  VA ECT BRASILIA   [posto]            ║
║    P11 ESCALA + TB + HORAS:                                            ║
║        · 12x36 ............... 47 (TB2) / 46 (TB1)                       ║
║        · 5x2 ou 6x1, 6h+ ..... 45 (TB2) / 44 (TB1)                       ║
║        · 5x2 ou 6x1, 4h ...... 43 (TB2) / 42 (TB1)                       ║
║    P12 Nada reconhecido ...... SEM benefício (entra no relatório)        ║
║                                                                          ║
║  EXTRAS (acumulativos — TODOS que baterem entram):                       ║
║    E1  Posto com 'IF -' ou '- IF' ........ 53  CB IF                     ║
║    E2  AMBIENTAL VALE + cozinha .......... 51  CB COZINHA VALE           ║
║    E3  Função MOTORISTA .................. 41  CB MOTORISTA              ║
║    E4  Posto VPORTS / MULTILIFT .......... 55  CF VPORTS                 ║
║    E5  Função GUARDA VIDA ................ 48  CF GUARDA VIDA            ║
║    E6  VALE SA + LAVADOR VEIC. PESADO .... 137 CB DESJEJUM VALE          ║
║    E7  VT (Grande Vitória-ES): CETURB → 58 · demais → 59 · fora → sem    ║
║                                                                          ║
║  OBS: o 136 (VA VALE) é sempre PRINCIPAL (regra P6, posto contém VALE).  ║
║  Pela regra de ouro, VA nunca é extra — por isso não há extra de VA.     ║
╚══════════════════════════════════════════════════════════════════════════╝
"""
from __future__ import annotations

import io
import re
from collections import Counter
from typing import Tuple, List, Dict

import pandas as pd

from .utils import normalizar_texto, calcular_carga_horaria, cidade_eh_grande_vitoria, eh_menor_de_idade
from .config import (
    RegrasConfig,
    PADRAO_POSTOS_TB2,
    PADRAO_POSTOS_VA_VPORTS,
    PADRAO_POSTOS_COFFEE_VPORTS,
    PADRAO_POSTOS_VT_CETURB,
)


# ══════════════════════════════════════════════════════════════════════════
# IDs DOS BENEFÍCIOS (EasyApp) — espelho da aba BENEFICIOS
# ══════════════════════════════════════════════════════════════════════════
ID_VA_SUPERVISOR       = 63    # VA ALIMENTACAO ADM 700
ID_VA_MOTOBOY          = 57    # VA ALIMENTACAO MOTOBOY
ID_VA_MOTORISTA        = 40    # VA ALIMENTACAO MOTORISTA
ID_VA_VIGILANTE        = 56    # VA ALIMENTACAO VIGILANTE VSP
ID_VA_VALE             = 136   # VA ALIMENTACAO VALE
ID_VA_COZINHA_VALE     = 50    # VA ALIMENTACAO I. AMB VALE (COZINHA)
ID_VA_VPORTS           = 54    # VA ALIMENTACAO VPORTS
ID_VA_IFE              = 52    # VA ALIMENTACAO IFE
ID_VA_ECT_BRASILIA     = 49    # VA ALIMENTACAO ECT BRASILIA
ID_VA_12X36_TB1        = 46
ID_VA_12X36_TB2        = 47
ID_VA_6H_TB1           = 44
ID_VA_6H_TB2           = 45
ID_VA_4H_TB1           = 42
ID_VA_4H_TB2           = 43
ID_CB_IF               = 53    # VA CESTA BASICA IF
ID_CB_COZINHA_VALE     = 51    # VA CESTA BASICA I. AMB VALE COZINHA
ID_CB_MOTORISTA        = 41    # VA CESTA BASICA MOTORISTA
ID_CF_VPORTS           = 55    # VA COFFEE BREAK VPORTS
ID_CF_GUARDA_VIDA      = 48    # VA COFFE BREAK GUARDA VIDA
ID_CB_DESJEJUM_VALE    = 137   # VA DESJEJUM VALE
ID_VT_CETURB           = 58    # VT TRANSPORTE - GVBUS 1 P/DIA
ID_VT_GRANDE_VITORIA   = 59    # VT TRANSPORTE - GVBUS 2 P/DIA

# ── Listas de postos (editáveis na aba LISTAS; estes são os padrões) ──────
POSTOS_TB2 = PADRAO_POSTOS_TB2
POSTOS_VA_VPORTS = PADRAO_POSTOS_VA_VPORTS
POSTOS_COFFEE_VPORTS = PADRAO_POSTOS_COFFEE_VPORTS
POSTOS_VT_CETURB = PADRAO_POSTOS_VT_CETURB

# ── Corte de carga horária para 6H / 4H ──────────────────────────────────
CORTE_HORAS_6H = 6.0

# ── Valores fixos do legado (o EasyApp espera exatamente estes) ───────────
VALORES_FIXOS = {
    ID_VA_VALE:          {"qtd_dias": 1, "valor_unico": 1106.70, "valor_total": 1106.70},
    ID_CB_DESJEJUM_VALE: {"qtd_dias": 1, "valor_unico": 154.49,  "valor_total": 154.49},
}


# ══════════════════════════════════════════════════════════════════════════
# TRATAMENTO DA PLANILHA DE BENEFÍCIOS (cadastro)
# ══════════════════════════════════════════════════════════════════════════
def tratar_base_beneficios(df: pd.DataFrame) -> pd.DataFrame:
    """Padroniza a planilha BENEFICIOS - RELACAO.xlsx para uso no motor."""
    df = df.loc[:, ~df.columns.str.contains("^Unnamed")].copy()
    df.columns = [c.strip().lower() for c in df.columns]
    colunas_esperadas = ["beneficio_id", "beneficio_nome", "sigla",
                         "valor_unico", "qtd_dias", "valor_total"]
    for col in colunas_esperadas:
        if col not in df.columns:
            df[col] = None
    df["sigla"] = df["sigla"].astype(str).str.strip().str.upper()
    df["beneficio_nome"] = df["beneficio_nome"].astype(str).str.strip().str.upper()
    # ⚠️ CRUCIAL: o Google Sheets devolve o ID como TEXTO ("56").
    # Sem esta conversão, a comparação "56" == 56 falha e NENHUM
    # benefício é encontrado (planilha sai vazia). Não remover!
    df["beneficio_id"] = pd.to_numeric(df["beneficio_id"], errors="coerce")
    # números podem vir com vírgula (15,5) vindos do Sheets — normaliza
    for col in ["valor_unico", "qtd_dias", "valor_total"]:
        df[col] = pd.to_numeric(
            df[col].astype(str).str.replace(",", ".", regex=False),
            errors="coerce",
        )
    df["valor_total"] = df["valor_total"].round(2)
    df["valor_unico"] = df["valor_unico"].round(2)
    # remove linhas sem ID (lixo no fim da planilha)
    df = df[df["beneficio_id"].notna()].copy()
    return df


def _buscar_por_id(df_benef: pd.DataFrame, beneficio_id: int) -> pd.DataFrame:
    """Atalho: devolve a linha da planilha com o ID informado."""
    return df_benef[df_benef["beneficio_id"] == beneficio_id]


def _contem_algum(texto: str, termos: list) -> bool:
    """True se `texto` (já normalizado) contém qualquer um dos `termos`."""
    return any(t in texto for t in termos)


def _comeca_com_ifes(posto: str) -> bool:
    """VA principal 52 (VA IFE): posto começa com 'IFES'.
    Ex.: IFES - VITORIA, IFES VENDA NOVA DO IMIGRANTE."""
    return bool(posto) and posto.startswith("IFES")


def _contem_if_separado(posto: str) -> bool:
    """Extra CB 53 (Cesta Básica IF): posto com 'IF -' ou '- IF'.
    Ex.: IF - CAMPUS PETROLINA, IF - CAMPUS FLORESTA, IF - CAMPUS OURICURI,
    IF - CAMPUS SERRA TALHADA."""
    if not posto:
        return False
    return "IF -" in posto or "- IF" in posto


# ══════════════════════════════════════════════════════════════════════════
# BENEFÍCIO PRINCIPAL (VA) — por ordem de prioridade
# ══════════════════════════════════════════════════════════════════════════
def encontrar_beneficios(
    row: pd.Series,
    df_benef: pd.DataFrame,
    regras: RegrasConfig | None = None,
) -> pd.DataFrame:
    """
    Devolve o benefício PRINCIPAL (VA) do colaborador.
    A primeira regra que bater vence — as demais não são avaliadas.
    """
    regras = regras or RegrasConfig()
    funcao  = normalizar_texto(row.get("nomefuncao"))
    escala  = normalizar_texto(row.get("tipoescala"))
    posto   = normalizar_texto(row.get("nomepostotrabalho"))
    horario = row.get("nomehorariotrabalho", "")
    carga   = calcular_carga_horaria(horario)

    # ── P0 · MENOR APRENDIZ (< 18 anos) — não recebe VA ─────────────────
    if eh_menor_de_idade(row.get("datanascimento")):
        return pd.DataFrame()

    # ── P1 · SUPERVISOR(A) → 63 ─────────────────────────────────────────
    if "SUPERVISOR" in funcao:
        return _buscar_por_id(df_benef, ID_VA_SUPERVISOR)

    # ── P2 · MOTOBOY → 57 ───────────────────────────────────────────────
    if "MOTOBOY" in funcao or "MOTO BOY" in funcao:
        return _buscar_por_id(df_benef, ID_VA_MOTOBOY)

    # ── P3 · MOTORISTA → 40 ─────────────────────────────────────────────
    if "MOTORISTA" in funcao:
        return _buscar_por_id(df_benef, ID_VA_MOTORISTA)

    # ── P4 · VIGILANTE → 56 ─────────────────────────────────────────────
    if _contem_algum(funcao, ["VIGILANTE", "VIGILANCIA"]):
        return _buscar_por_id(df_benef, ID_VA_VIGILANTE)

    # ── P5 · COZINHEIRO / AJUDANTE DE COZINHA do AMBIENTAL VALE → 50 ────
    #    Regra mais específica: o pessoal da cozinha do Instituto Ambiental
    #    Vale recebe o VA próprio da cozinha (50), NÃO o 136 genérico.
    #    (CONEXAO + anotação do setor: "COZINHEIRO - Vale (VA cozinha + CB)")
    if "AMBIENTAL VALE" in posto and _contem_algum(funcao, ["COZINHEIRO", "AJUDANTE DE COZINHA"]):
        return _buscar_por_id(df_benef, ID_VA_COZINHA_VALE)

    # ── P6 · POSTO VALE SA → 136 ────────────────────────────────────────
    if "VALE" in posto:
        return _buscar_por_id(df_benef, ID_VA_VALE)

    # ── P7 / P8 · POSTO VPORTS ou MULTILIFT → 54 ────────────────────────
    if _contem_algum(posto, regras.postos_va_vports):
        return _buscar_por_id(df_benef, ID_VA_VPORTS)

    # ── P9 · POSTO que começa com "IFES" → 52 ───────────────────────────
    if _comeca_com_ifes(posto):
        return _buscar_por_id(df_benef, ID_VA_IFE)

    # ── P10 · POSTO contém BRASILIA → 49 ────────────────────────────────
    if "BRASILIA" in posto:
        return _buscar_por_id(df_benef, ID_VA_ECT_BRASILIA)

    # ── P11 · ESCALA + TB + HORAS ───────────────────────────────────────
    filtro_tb = "TB2" if _contem_algum(posto, regras.postos_tb2) else "TB1"

    if "12X36" in escala:
        alvo = ID_VA_12X36_TB2 if filtro_tb == "TB2" else ID_VA_12X36_TB1
        return _buscar_por_id(df_benef, alvo)

    if _contem_algum(escala, ["5X2", "6X1"]):
        if carga >= CORTE_HORAS_6H:
            alvo = ID_VA_6H_TB2 if filtro_tb == "TB2" else ID_VA_6H_TB1
        else:
            alvo = ID_VA_4H_TB2 if filtro_tb == "TB2" else ID_VA_4H_TB1
        return _buscar_por_id(df_benef, alvo)

    # ── P12 · NADA RECONHECIDO → sem benefício principal ────────────────
    return pd.DataFrame()


# ══════════════════════════════════════════════════════════════════════════
# BENEFÍCIOS EXTRAS (CB / CF / VT) — acumulativos
# ══════════════════════════════════════════════════════════════════════════
def encontrar_extras(
    row: pd.Series,
    df_benef: pd.DataFrame,
    regras: RegrasConfig | None = None,
) -> pd.DataFrame:
    """
    Devolve TODOS os extras (CB/CF/VT) do colaborador.
    As regras são independentes e acumulam entre si.
    """
    regras = regras or RegrasConfig()
    funcao = normalizar_texto(row.get("nomefuncao"))
    posto  = normalizar_texto(row.get("nomepostotrabalho"))

    # Cidade de residência: prefere a do EasyApp (cidadeparceiro);
    # cai para cidadetrabalho se a primeira estiver vazia.
    cidade = row.get("cidadeparceiro")
    if pd.isna(cidade) or str(cidade).strip() == "":
        cidade = row.get("cidadetrabalho")

    extras_ids = set()

    # ── E1 · CB — posto com "IF -" ou "- IF" → 53 ───────────────────────
    if _contem_if_separado(posto):
        extras_ids.add(ID_CB_IF)

    # ── E2 · CB — cozinha do AMBIENTAL VALE → 51 ────────────────────────
    if "AMBIENTAL VALE" in posto and _contem_algum(funcao, ["COZINHEIRO", "AJUDANTE DE COZINHA"]):
        extras_ids.add(ID_CB_COZINHA_VALE)

    # ── E3 · CB — função MOTORISTA → 41 ─────────────────────────────────
    if "MOTORISTA" in funcao:
        extras_ids.add(ID_CB_MOTORISTA)

    # ── E4 · CF — posto VPORTS / MULTILIFT → 55 ─────────────────────────
    if _contem_algum(posto, regras.postos_coffee_vports):
        extras_ids.add(ID_CF_VPORTS)

    # ── E5 · CF — função GUARDA-VIDA → 48 ───────────────────────────────
    if "GUARDA VIDA" in funcao:
        extras_ids.add(ID_CF_GUARDA_VIDA)

    # ── E6 · CB — lavador de veículos pesado do VALE → 137 ──────────────
    if "VALE" in posto and "LAVADOR DE VEICULOS PESADO" in funcao:
        extras_ids.add(ID_CB_DESJEJUM_VALE)

    # ── E7 · VALE-TRANSPORTE (somente Grande Vitória - ES) ──────────────
    #    (não existe extra de VA: pela regra de ouro, VA é sempre principal)
    if cidade_eh_grande_vitoria(cidade, regras.cidades_grande_vitoria):
        if _contem_algum(posto, regras.postos_vt_ceturb):
            extras_ids.add(ID_VT_CETURB)
        else:
            extras_ids.add(ID_VT_GRANDE_VITORIA)

    if not extras_ids:
        return pd.DataFrame()
    return df_benef[df_benef["beneficio_id"].isin(extras_ids)].copy()


# ══════════════════════════════════════════════════════════════════════════
# PROCESSADOR PRINCIPAL — gera a planilha "Importar Beneficios"
# ══════════════════════════════════════════════════════════════════════════
def processar_beneficios(
    df_func: pd.DataFrame,
    beneficios_relacao_bytes: bytes | None = None,
    ultimo_id_easyapp: int | None = None,
    usar_id_sequencial: bool = False,
    regras: RegrasConfig | None = None,
    df_benef_pronto: pd.DataFrame | None = None,
) -> Tuple[pd.DataFrame, List[Dict]]:
    """
    Gera o DataFrame da planilha "Importar Beneficios".

    - `df_func`: DataFrame da Importar Layout.
    - `beneficios_relacao_bytes`: bytes do arquivo BENEFICIOS - RELACAO.xlsx.
    - `ultimo_id_easyapp` / `usar_id_sequencial`: IDs sequenciais a partir
      do último ID do EasyApp; senão, colaborador_id = matrícula do eSocial.
    - `regras`: listas configuráveis (Google Sheets / aba LISTAS).
    - `df_benef_pronto`: DataFrame de benefícios já carregado (Google Sheets);
      quando informado, `beneficios_relacao_bytes` é ignorado.

    Retorna: (DataFrame de saída, relatório de erros).
    """
    regras = regras or RegrasConfig()

    if df_benef_pronto is not None:
        df_benef = tratar_base_beneficios(df_benef_pronto)
    elif beneficios_relacao_bytes is not None:
        df_benef = tratar_base_beneficios(pd.read_excel(io.BytesIO(beneficios_relacao_bytes)))
    else:
        raise ValueError("Informe a planilha de benefícios (bytes) ou df_benef_pronto.")

    # garante que df_func tem as colunas usadas pelas regras
    for c in ["matricula", "nome", "nomefuncao", "nomepostotrabalho",
              "tipoescala", "nomehorariotrabalho", "datanascimento",
              "cidadeparceiro", "cidadetrabalho"]:
        if c not in df_func.columns:
            df_func[c] = ""

    df_func = df_func.copy().reset_index(drop=True)

    # ────────── colaborador_id sequencial (opcional) ──────────
    if usar_id_sequencial and ultimo_id_easyapp is not None:
        df_func["__colaborador_id"] = [
            int(ultimo_id_easyapp) + i + 1 for i in range(len(df_func))
        ]
    else:
        df_func["__colaborador_id"] = df_func["matricula"]

    # ────────── aplica as regras linha a linha ──────────
    linhas: List[pd.Series] = []
    for _, row in df_func.iterrows():
        principais = encontrar_beneficios(row, df_benef, regras)
        extras     = encontrar_extras(row, df_benef, regras)
        encontrados = pd.concat([principais, extras], ignore_index=True)
        encontrados = encontrados.drop_duplicates(subset=["beneficio_id"])

        if encontrados.empty:
            # colaborador sem nenhum benefício → linha em branco + relatório
            nova = row.copy()
            for c in ["beneficio_id", "beneficio_nome", "qtd_dias",
                      "valor_unico", "valor_total", "sigla"]:
                nova[c] = None
            linhas.append(nova)
        else:
            for _, b in encontrados.iterrows():
                if pd.isna(b["beneficio_id"]):
                    continue
                nova = row.copy()
                bid = int(b["beneficio_id"])
                nova["beneficio_id"] = bid
                nova["beneficio_nome"] = b["beneficio_nome"]
                nova["sigla"] = b["sigla"]
                if bid in VALORES_FIXOS:      # 136 e 137: valores fixos do legado
                    fixo = VALORES_FIXOS[bid]
                    nova["qtd_dias"], nova["valor_unico"], nova["valor_total"] = (
                        fixo["qtd_dias"], fixo["valor_unico"], fixo["valor_total"]
                    )
                else:
                    nova["qtd_dias"]    = b["qtd_dias"]
                    nova["valor_unico"] = b["valor_unico"]
                    nova["valor_total"] = b["valor_total"]
                linhas.append(nova)

    df_final = pd.DataFrame(linhas)
    df_final["erro_beneficio"] = df_final["beneficio_id"].isna()

    # ────────── layout de saída (novo modelo LAYOUT_DetalheBeneficio) ──────────
    # Colunas: id (vazio) | parceiro_id (colaborador) | beneficio_id
    df_saida = pd.DataFrame()
    df_saida["id"]           = [""] * len(df_final)
    df_saida["parceiro_id"]  = df_final["__colaborador_id"].values
    df_saida["beneficio_id"] = df_final["beneficio_id"].values

    # ────────── relatório de erros (sem benefício identificado) ──────────
    erros = df_final[df_final["erro_beneficio"]]
    relatorio = [
        {
            "matricula": r.get("matricula"),
            "nome": r.get("nome"),
            "funcao": r.get("nomefuncao"),
            "posto": r.get("nomepostotrabalho"),
        }
        for _, r in erros.iterrows()
    ]

    return df_saida, relatorio


# ══════════════════════════════════════════════════════════════════════════
# CORES NA PLANILHA "IMPORTAR BENEFICIOS"  (regra revisada em set/2026)
# ------------------------------------------------------------------------
# A cor é aplicada na célula do `parceiro_id`, olhando APENAS os
# benefícios de sigla VA / CB / CF (o VT é tratado à parte). Como o novo
# layout não tem a coluna "sigla", o VT é identificado pelo beneficio_id
# (58 e 59 = VT TRANSPORTE GVBUS):
#
#   🔵 AZUL     → colaborador com MAIS DE UM benefício VA/CB/CF
#   🔴 VERMELHO → colaborador com APENAS UM benefício VA/CB/CF
#   🟢 VERDE    → linhas de VT (vale-transporte), independente da contagem
# ══════════════════════════════════════════════════════════════════════════
IDS_VT = {ID_VT_CETURB, ID_VT_GRANDE_VITORIA}   # 58 e 59


def aplicar_cores_colaborador_id(xlsx_bytes: bytes) -> bytes:
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    wb = load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active

    header = {cell.value: i for i, cell in enumerate(ws[1], start=1)}
    # novo layout usa "parceiro_id"; mantém compatibilidade com o antigo
    col_id   = header.get("parceiro_id") or header.get("colaborador_id")
    col_ben  = header.get("beneficio_id")
    col_sigla = header.get("sigla")   # layout antigo (se existir)
    if col_id is None:
        return xlsx_bytes

    def _eh_vt(row: int) -> bool:
        """True se a linha é um VT (pela sigla ou, no layout novo, pelo ID)."""
        if col_sigla:
            sigla = ws.cell(row=row, column=col_sigla).value
            return str(sigla).strip().upper() == "VT" if sigla else False
        if col_ben:
            try:
                return int(float(ws.cell(row=row, column=col_ben).value)) in IDS_VT
            except (TypeError, ValueError):
                return False
        return False

    fill_red   = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    fill_blue  = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
    fill_green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    # 1) conta, por colaborador, quantos benefícios VA/CB/CF ele tem
    #    (linhas VT NÃO entram na contagem)
    contagem: Counter = Counter()
    for row in range(2, ws.max_row + 1):
        v = ws.cell(row=row, column=col_id).value
        chave = "" if v is None else str(v).strip()
        if chave and not _eh_vt(row):
            contagem[chave] += 1

    # 2) pinta: VT sempre verde; demais → azul (2+ VA/CB/CF) ou vermelho (1)
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=col_id)
        v = "" if cell.value is None else str(cell.value).strip()
        if _eh_vt(row):
            cell.fill = fill_green
        else:
            cell.fill = fill_blue if contagem[v] > 1 else fill_red

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
