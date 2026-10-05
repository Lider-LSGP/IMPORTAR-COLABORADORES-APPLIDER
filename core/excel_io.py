# -*- coding: utf-8 -*-
"""
Leitura inteligente de planilhas e conversão XLSX -> CSV preservando texto.

A conversão preserva zeros à esquerda (CPF, PIS, matrícula, CEP...) exatamente
como no script original converter-csv.py / app_importador.py.
"""
import csv
import io
import re
from datetime import datetime, date

import pandas as pd
from openpyxl import load_workbook


def _to_bytes(file):
    """Aceita UploadedFile do Streamlit, BytesIO, path (str/Path) ou bytes."""
    if file is None:
        return None
    if isinstance(file, (bytes, bytearray)):
        return bytes(file)
    if hasattr(file, "getvalue"):
        return file.getvalue()
    if hasattr(file, "read"):
        pos = None
        try:
            pos = file.tell()
        except Exception:
            pass
        data = file.read()
        if pos is not None:
            try:
                file.seek(pos)
            except Exception:
                pass
        return data
    with open(file, "rb") as f:
        return f.read()


def read_excel_smart(file, dtype=None, sheet_name=0):
    """
    Lê .xls/.xlsx reais e também '.xls' que na verdade são HTML
    (exportações comuns de sistemas como Domínio / EasyApp).
    """
    data = _to_bytes(file)
    head = data[:2048].lstrip().lower()
    is_html = head.startswith(b"<") or b"<html" in head or b"<table" in head

    if is_html:
        tables = pd.read_html(io.BytesIO(data), header=0)
        df = tables[0]
        if dtype is str:
            df = df.astype(str)
        return df

    try:
        return pd.read_excel(io.BytesIO(data), dtype=dtype, sheet_name=sheet_name)
    except Exception:
        tables = pd.read_html(io.BytesIO(data), header=0)
        df = tables[0]
        if dtype is str:
            df = df.astype(str)
        return df


def valor_exato_da_celula(cell):
    """Preserva o valor exatamente como aparece no Excel (zeros à esquerda etc.)."""
    valor = cell.value
    if valor is None:
        return ""
    if isinstance(valor, str):
        return valor
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%Y-%m-%d")

    formato = (cell.number_format or "").split(";")[0].strip()
    if re.fullmatch(r"0+", formato):
        try:
            return f"{int(valor):0{len(formato)}d}"
        except Exception:
            return str(valor)

    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


def xlsx_bytes_to_csv_bytes(xlsx_bytes):
    """Converte o conteúdo de um .xlsx (bytes) para CSV (bytes utf-8) COM cabeçalho."""
    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    ws = wb.active
    out = io.StringIO()
    writer = csv.writer(out)
    for row in ws.iter_rows():
        writer.writerow([valor_exato_da_celula(c) for c in row])
    return out.getvalue().encode("utf-8")
