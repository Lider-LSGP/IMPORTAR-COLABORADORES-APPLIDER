# -*- coding: utf-8 -*-
"""
Exportadores XLSX com formatação visual (cores) via openpyxl.

- df_to_xlsx_bytes(..., red_cells=...)  -> pinta de vermelho células específicas
- df_to_xlsx_bytes(..., red_all=True)   -> pinta a planilha inteira de vermelho
- beneficios_xlsx_bytes(...)            -> regra original: azul = colaborador_id
                                           repetido, vermelho = único
"""
import io
from collections import Counter

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

FILL_RED = PatternFill(start_color="FF4C4C", end_color="FF4C4C", fill_type="solid")
FILL_BLUE = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")


def df_to_xlsx_bytes(df, red_cells=None, red_all=False):
    """
    red_cells: lista de tuplas (indice_linha_zero_based, nome_coluna)
    red_all:   pinta TODAS as células de dados de vermelho
    """
    buf = io.BytesIO()
    df.to_excel(buf, index=False)
    buf.seek(0)

    wb = load_workbook(buf)
    ws = wb.active

    if red_all:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.fill = FILL_RED
    elif red_cells:
        colmap = {c.value: i for i, c in enumerate(ws[1], start=1)}
        for row_idx, col_name in red_cells:
            ci = colmap.get(col_name)
            if ci:
                ws.cell(row=row_idx + 2, column=ci).fill = FILL_RED

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def beneficios_xlsx_bytes(df_saida):
    """
    Gera o Importar Beneficios.xlsx com a regra visual original:
      - colaborador_id REPETIDO (mais de um benefício) -> AZUL
      - colaborador_id ÚNICO -> VERMELHO
    """
    buf = io.BytesIO()
    df_saida.to_excel(buf, index=False)
    buf.seek(0)

    wb = load_workbook(buf)
    ws = wb.active
    colmap = {c.value: i for i, c in enumerate(ws[1], start=1)}
    ci = colmap.get("colaborador_id")

    if ci:
        valores = []
        for r in range(2, ws.max_row + 1):
            val = ws.cell(row=r, column=ci).value
            valores.append("" if val is None else str(val).strip())
        contagem = Counter(valores)
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=ci)
            val = "" if cell.value is None else str(cell.value).strip()
            cell.fill = FILL_BLUE if contagem[val] > 1 else FILL_RED

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
